"""Generate small real fixtures for the XLSX/DOCX/PPTX/PDF connectors.

Writes into tests/fixtures/. Requires the optional connector libraries; used
by the test suite (which skips gracefully when a library or fixture is absent)
and for manual demos. Run: python scripts/make_fixtures.py
"""

from __future__ import annotations

import os

FIX = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures")
os.makedirs(FIX, exist_ok=True)


def make_xlsx() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cabinet Tracker"
    # Merged title cell across row 1; subtitle row 2; real header row 3.
    ws["A1"] = "मंत्रिमंडळ निर्णय ट्रॅकर 2025"
    ws.merge_cells("A1:E1")
    ws["A2"] = "गोपनीय"
    headers = ["अ.क्र.", "दिनांक", "विभाग", "प्रवर्ग", "तपशील"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=h)
    rows = [
        [1, "2025-03-12", "सामान्य प्रशासन", "Cabinet Decisions",
         "SAMAGRA एकात्मिक लाभार्थी डेटाबेस स्वीकारण्याचा निर्णय"],
        [2, "2025-04-02", "ऊर्जा", "Cabinet Decisions",
         "सौर वीज धोरण 2025 ला मंजुरी"],
    ]
    for r, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(os.path.join(FIX, "cabinet_tracker.xlsx"))


def make_docx() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Minutes of PS Meeting", level=1)
    d.add_paragraph("Date: 5 May 2025. Chaired by the Principal Secretary.")
    d.add_heading("Water supply review", level=2)
    d.add_paragraph(
        "The Water Supply department presented the पाणीपुरवठा mission status "
        "for key districts."
    )
    d.save(os.path.join(FIX, "ps_meeting_minutes.docx"))


def make_pptx() -> None:
    import pptx
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quantum Mission Overview"
    body = slide.placeholders[1]
    body.text = "Quantum technology park proposed at Pune (क्वांटम तंत्रज्ञान उद्यान)."
    slide.notes_slide.notes_text_frame.text = (
        "Speaker notes: outlay to be finalised by the Industries department."
    )
    prs.save(os.path.join(FIX, "quantum_deck.pptx"))


def make_pdf() -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        "REPORT: Semiconductor manufacturing investment\n"
        "Land acquisition near Dhule-Nashik is under way.",
        fontsize=11,
    )
    doc.save(os.path.join(FIX, "semiconductor_report.pdf"))
    doc.close()


def main() -> None:
    make_xlsx()
    make_docx()
    make_pptx()
    make_pdf()
    print("fixtures written to", os.path.abspath(FIX))


if __name__ == "__main__":
    main()
