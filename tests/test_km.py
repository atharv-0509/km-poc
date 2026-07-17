"""Unit + integration tests for the KM PoC. Run: python -m unittest -v

Uses only the standard library and the default hashing embedder, so the
whole suite runs offline with no third-party dependencies.
"""

from __future__ import annotations

import os
import tempfile
import unittest

from km.config import Config
from km.embeddings import HashingEmbedder, cosine
from km.ingest import ingest_path
from km.ingest.spreadsheet import detect_header, _dedupe_headers
from km.lang import detect_language, expand_terms
from km.pipeline import build_index, search
from km.schema import Record
from km.tagging import parse_date, infer_category, tag


def _has(mod: str) -> bool:
    import importlib.util
    return importlib.util.find_spec(mod) is not None


class TestSchema(unittest.TestCase):
    def test_stable_id_and_provenance(self):
        r = Record(title="T", raw_text="x", source_file="f.csv",
                   sheet="s", row_or_page="row 4")
        self.assertEqual(r.id, r.compute_id())
        self.assertIn("f.csv", r.provenance)
        self.assertIn("row 4", r.provenance)

    def test_roundtrip(self):
        r = Record(title="T", raw_text="body", source_file="f", category="War Room")
        r2 = Record.from_dict(r.to_dict())
        self.assertEqual(r2.id, r.id)
        self.assertEqual(r2.category, "War Room")


class TestTagging(unittest.TestCase):
    def test_dates(self):
        self.assertEqual(parse_date("dated 14-06-2025 at"), "2025-06-14")
        self.assertEqual(parse_date("on 2025/03/04"), "2025-03-04")
        self.assertEqual(parse_date("19 April 2025"), "2025-04-19")
        self.assertIsNone(parse_date("no date here"))

    def test_language_detection(self):
        self.assertEqual(detect_language("Letter to the President"), "eng")
        self.assertEqual(detect_language("राष्ट्रपती महोदयांना निमंत्रण आहे"), "mar")
        self.assertEqual(detect_language("Subject: क्वांटम मिशन and quantum park review"), "mixed")

    def test_category_from_text(self):
        self.assertEqual(infer_category("cabinet approved the scheme"), "Cabinet Decisions")
        self.assertEqual(infer_category("वॉर रूम आढावा"), "War Room")

    def test_tag_fills_metadata(self):
        r = Record(title="Letter to Hon'ble President", raw_text="dated 2025-02-11 Mumbai",
                   source_file="vip.csv")
        tag(r)
        self.assertEqual(r.date, "2025-02-11")
        self.assertIn("Mumbai", r.location)
        self.assertTrue(any("President" in e for e in r.entities))


class TestCrossLingual(unittest.TestCase):
    def test_expansion_bridges_scripts(self):
        aliases = expand_terms("letter to the president")
        joined = " ".join(aliases)
        self.assertIn("राष्ट्रपती", joined)
        self.assertIn("पत्र", joined)

    def test_semiconductor(self):
        self.assertIn("सेमीकंडक्टर", " ".join(expand_terms("semiconductor")))


class TestHeaderDetection(unittest.TestCase):
    def test_header_below_row_one(self):
        rows = [
            ["Announcement Tracker 2025", "", "", ""],  # merged title row
            ["confidential", "", "", ""],               # subtitle
            ["Sr", "Date", "Dept", "Detail"],           # real header (idx 2)
            ["1", "2025-01-01", "Home", "something"],
            ["2", "2025-02-01", "Finance", "another"],
        ]
        self.assertEqual(detect_header(rows), 2)

    def test_dedupe_headers(self):
        self.assertEqual(
            _dedupe_headers(["Status", "Status", "", "Date"]),
            ["Status", "Status_1", "col3", "Date"],
        )


class TestEmbedder(unittest.TestCase):
    def test_normalised_and_self_similar(self):
        emb = HashingEmbedder(dim=256)
        v = emb.embed("quantum technology park")
        self.assertAlmostEqual(cosine(v, v), 1.0, places=5)
        near = cosine(v, emb.embed("quantum technology centre"))
        far = cosine(v, emb.embed("water supply metro corridor"))
        self.assertGreater(near, far)


class TestEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.cfg = Config(
            db_path=os.path.join(cls.tmp, "idx.sqlite3"),
            data_dir="data", embedder="hashing",
        )
        cls.summary = build_index(cls.cfg)

    def test_indexed_expected_counts(self):
        self.assertGreaterEqual(self.summary["indexed"], 25)
        self.assertIn("scan", self.summary["by_source_type"])
        self.assertIn("letter", self.summary["by_source_type"])

    def _titles(self, query, **kw):
        res = search(self.cfg, query, **kw)
        return [h.record.title for h in res.hits], res

    def test_quantum_query(self):
        _titles, res = self._titles("Show me everything on Quantum")
        files = {h.record.source_file for h in res.hits}
        self.assertIn("scan_quantum_note.txt", files)
        self.assertIn("announcements_tracker.csv", files)

    def test_cross_lingual_president(self):
        # English query must retrieve Marathi/Hindi राष्ट्रपती records.
        _t, res = self._titles("Letters to the President in 2025",
                               filters={"date": ["2025-01-01", "2025-12-31"]})
        langs = {h.record.language for h in res.hits}
        # At least one non-English record surfaced from an English query.
        self.assertTrue({"mar", "hin", "mixed"} & langs)
        # The 2024 President letter must be excluded by the date filter.
        for h in res.hits:
            if h.record.date:
                self.assertTrue("2025" in h.record.date)

    def test_semiconductor_meeting(self):
        _t, res = self._titles("What meetings on semiconductor manufacturing?")
        files = {h.record.source_file for h in res.hits}
        self.assertIn("scan_cm_meeting_semiconductor.txt", files)

    def test_category_filter(self):
        _t, res = self._titles("पाणीपुरवठा", filters={"category": "War Room"})
        self.assertTrue(res.hits)
        for h in res.hits:
            self.assertEqual(h.record.category, "War Room")


class TestFieldAwareRanking(unittest.TestCase):
    """Recipient/subject fields must outrank incidental body mentions, and
    common stopwords must not dilute keyword ranking."""

    def setUp(self):
        from km.embeddings import HashingEmbedder
        from km.store import Store
        self.tmp = tempfile.mkdtemp()
        self.store = Store(os.path.join(self.tmp, "f.sqlite3"), HashingEmbedder())
        # A: addressed TO the President. B: merely mentions 'president' in body.
        a = Record(title="Invitation to grace the ceremony", raw_text="body text",
                   source_file="letters.xlsx", sheet="VIP", row_or_page="row 2",
                   source_type="letter", key_fields="H.E. President of India")
        b = Record(title="Budget note", raw_text="the vice president of the club "
                   "attended along with the president of the society",
                   source_file="letters.xlsx", sheet="VIP", row_or_page="row 3",
                   source_type="letter", key_fields="Shri Ramesh Kumar")
        # Filler so 'president' is a rare term (BM25 IDF is meaningless in a
        # 2-doc corpus — it goes negative when a term is in most documents).
        filler = [
            Record(title=f"Routine letter {i}", raw_text=f"matter number {i} noted",
                   source_file="letters.xlsx", sheet="VIP", row_or_page=f"row {10+i}",
                   source_type="letter", key_fields=f"Shri Official {i}")
            for i in range(12)
        ]
        self.store.add_many([tag(a), tag(b), *[tag(f) for f in filler]])
        self.a_id = a.id

    def tearDown(self):
        self.store.close()

    def test_recipient_outranks_body_mention(self):
        hits = self.store.keyword_search("letters to the President", limit=5)
        self.assertTrue(hits)
        self.assertEqual(hits[0][0], self.a_id)

    def test_stopwords_dropped(self):
        from km.lang import content_terms
        terms = content_terms("Show me all the letters to the President in 2025")
        self.assertNotIn("the", terms)
        self.assertNotIn("2025", terms)   # bare year handled by date filter
        self.assertIn("president", terms)


class TestFormatConnectors(unittest.TestCase):
    """Real-file connectors, each skipped unless its optional lib is present."""

    def _tag_all(self, path):
        return [tag(r) for r in ingest_path(path)]

    @unittest.skipUnless(_has("openpyxl"), "openpyxl not installed")
    def test_xlsx_header_and_merged_cells(self):
        import openpyxl
        tmp = tempfile.mkdtemp()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Cabinet Tracker 2025"          # merged title
        ws.merge_cells("A1:D1")
        ws["A2"] = "confidential"                   # subtitle
        for c, h in enumerate(["Sr", "Date", "प्रवर्ग", "तपशील"], start=1):
            ws.cell(row=3, column=c, value=h)       # real header on row 3
        ws.append([1, "2025-03-12", "Cabinet Decisions", "SAMAGRA database adopted"])
        p = os.path.join(tmp, "cabinet_tracker.xlsx")
        wb.save(p)

        recs = self._tag_all(p)
        self.assertEqual(len(recs), 1)
        r = recs[0]
        self.assertEqual(r.category, "Cabinet Decisions")   # from explicit column
        self.assertEqual(r.row_or_page, "row 4")
        self.assertIn("SAMAGRA", r.raw_text)

    @unittest.skipUnless(_has("docx"), "python-docx not installed")
    def test_docx_sections(self):
        import docx
        tmp = tempfile.mkdtemp()
        d = docx.Document()
        d.add_heading("Minutes of PS Meeting", level=1)
        d.add_paragraph("Chaired by the Principal Secretary.")
        d.add_heading("Water supply review", level=2)
        d.add_paragraph("पाणीपुरवठा mission status presented.")
        p = os.path.join(tmp, "ps_meeting_minutes.docx")
        d.save(p)

        recs = self._tag_all(p)
        self.assertGreaterEqual(len(recs), 2)
        self.assertTrue(all(r.source_type == "minutes" for r in recs))

    @unittest.skipUnless(_has("pptx"), "python-pptx not installed")
    def test_pptx_slide_and_notes(self):
        import pptx
        tmp = tempfile.mkdtemp()
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Quantum Mission Overview"
        slide.placeholders[1].text = "Quantum park at Pune"
        slide.notes_slide.notes_text_frame.text = "outlay pending"
        p = os.path.join(tmp, "quantum_deck.pptx")
        prs.save(p)

        recs = self._tag_all(p)
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0].source_type, "deck")
        self.assertIn("outlay", recs[0].raw_text)   # notes captured

    @unittest.skipUnless(_has("fitz"), "pymupdf not installed")
    def test_pdf_text_layer(self):
        import fitz
        tmp = tempfile.mkdtemp()
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "REPORT: Semiconductor manufacturing", fontsize=11)
        p = os.path.join(tmp, "semiconductor_report.pdf")
        doc.save(p)
        doc.close()

        recs = self._tag_all(p)
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0].row_or_page, "page 1")
        self.assertIn("Semiconductor", recs[0].raw_text)


if __name__ == "__main__":
    unittest.main()
